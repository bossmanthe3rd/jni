#!/usr/bin/env python3
"""Turn India Post's next-day parcel list into the site's pincode data.

The owner's spreadsheet carries two sheets. "24 Speed Post and 48 Speed Post"
mixes next-day and two-day areas and writes Delhi and Chennai as ranges, so it
cannot answer "will this arrive the next day". "24 Speed Post Parcel" can:
one row per pincode, six metros, and parcels are what the orders ship as.
That sheet is the whole source of truth for the PIN check.

Writes src/data/oneDayPincodes.json:

    { "source": ..., "cities": ["Bengaluru", ...], "pins": { "560001": 0, ... } }

`pins` maps each pincode to its city's index in `cities`, which keeps the file
small (it is fetched only when someone first uses a PIN field). Rows that are
not a clean six-digit pincode are reported and skipped, never guessed at.

Run `python tools/build-pincodes.py [path/to/list.xlsx]` from the repo root.
With no argument it reads the spreadsheet the owner dropped in the repo root.
"""

import json
import os
import re
import sys

import openpyxl

DEFAULT_SRC = "24 and 48 speed post and parcel pincode list 6 cities.xlsx"
SHEET = "24 Speed Post Parcel"
OUT = os.path.join("src", "data", "oneDayPincodes.json")

# The sheet spells two cities two ways ("Delhi" / "New Delhi" across sheets,
# shouty "HYDERABAD"); one display name each.
CITY_NAMES = {
    "delhi": "Delhi",
    "new delhi": "Delhi",
    "mumbai": "Mumbai",
    "kolkata": "Kolkata",
    "chennai": "Chennai",
    "hyderabad": "Hyderabad",
    "bengaluru": "Bengaluru",
    "bangalore": "Bengaluru",
}

PIN = re.compile(r"[1-9]\d{5}")


def main():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    os.chdir(root)
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SRC

    book = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if SHEET not in book.sheetnames:
        sys.exit("No sheet named %r in %s (found: %s)" % (SHEET, src, ", ".join(book.sheetnames)))

    cities, pins, skipped = [], {}, []
    # Row 1 is the sheet's title, row 2 its column headings.
    for row in list(book[SHEET].iter_rows(values_only=True))[2:]:
        if not any(row):
            continue
        city_raw, pin_raw = (row[1] or ""), row[3]
        pin = str(int(pin_raw)) if isinstance(pin_raw, (int, float)) else str(pin_raw or "").strip()
        city = CITY_NAMES.get(str(city_raw).strip().lower())
        if not PIN.fullmatch(pin) or not city:
            skipped.append(row)
            continue
        if city not in cities:
            cities.append(city)
        pins[pin] = cities.index(city)

    data = {"source": "%s / %s" % (os.path.basename(src), SHEET), "cities": cities, "pins": dict(sorted(pins.items()))}
    with open(OUT, "w", encoding="utf-8", newline="\n") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
        f.write("\n")

    print("  %d pincodes across %d cities -> %s (%.1f KB)" % (len(pins), len(cities), OUT, os.path.getsize(OUT) / 1024))
    for c in cities:
        print("    %-10s %d" % (c, sum(1 for i in pins.values() if cities[i] == c)))
    for row in skipped:
        print("  skipped:", row)


if __name__ == "__main__":
    main()
